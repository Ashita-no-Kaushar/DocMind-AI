import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from utils.format_ingestion import (
    CAPABILITY_CATEGORIES,
    PARSER_CAPABILITIES,
    SUPPORTED_EXTENSIONS,
    build_file_extractors,
)
from utils.helpers import ALLOWED_UPLOAD_EXTENSIONS
from utils.llama_index import load_documents


class FormatCapabilityTests(unittest.TestCase):
    def test_registry_matches_existing_25_extension_contract(self):
        expected = {
            "csv",
            "doc",
            "docx",
            "eml",
            "epub",
            "htm",
            "html",
            "ipynb",
            "json",
            "jsonl",
            "markdown",
            "mbox",
            "md",
            "mhtml",
            "msg",
            "odt",
            "pdf",
            "ppt",
            "pptx",
            "rtf",
            "tsv",
            "txt",
            "xls",
            "xlsx",
            "xml",
        }

        self.assertEqual(set(SUPPORTED_EXTENSIONS), expected)
        self.assertEqual(set(PARSER_CAPABILITIES), expected)
        self.assertEqual(len(SUPPORTED_EXTENSIONS), 25)
        self.assertEqual(len(PARSER_CAPABILITIES), 25)
        self.assertEqual(
            ALLOWED_UPLOAD_EXTENSIONS, {f".{extension}" for extension in expected}
        )
        self.assertTrue(
            all(capability.category in CAPABILITY_CATEGORIES for capability in PARSER_CAPABILITIES.values())
        )

    def test_legacy_formats_are_explicitly_unavailable_without_conversion(self):
        self.assertEqual(
            PARSER_CAPABILITIES["doc"].category,
            "unsupported_without_converter",
        )
        self.assertEqual(
            PARSER_CAPABILITIES["ppt"].category,
            "unsupported_without_converter",
        )

    def test_reader_mapping_covers_every_accepted_extension(self):
        readers = build_file_extractors()

        self.assertEqual(set(readers), {f".{ext}" for ext in SUPPORTED_EXTENSIONS})
        self.assertEqual(type(readers[".ppt"]).__name__, "LegacyFormatReader")
        self.assertEqual(type(readers[".pptx"]).__name__, "PptxReader")
        self.assertEqual(type(readers[".html"]).__name__, "WholeDocumentHTMLReader")
        self.assertEqual(type(readers[".jsonl"]).__name__, "JSONLReader")


class TextFormatExtractionTests(unittest.TestCase):
    def _load(self, directory, name):
        return load_documents(
            str(directory),
            input_files=[str(Path(directory) / name)],
            return_report=True,
        )

    def test_jsonl_reads_each_record_line_by_line(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "records.jsonl"
            path.write_text(
                '{"name":"first-jsonl-sentinel","count":1}\n'
                '{"name":"second-jsonl-sentinel","count":2}\n',
                encoding="utf-8",
            )

            documents, report = self._load(tmpdir, path.name)

        self.assertEqual(len(documents), 2)
        combined = "\n".join(document.text for document in documents)
        self.assertIn("first-jsonl-sentinel", combined)
        self.assertIn("second-jsonl-sentinel", combined)
        self.assertIn("Record 1:", combined)
        self.assertIn("Record 2:", combined)
        self.assertEqual(report[0]["status"], "loaded")
        self.assertEqual(report[0]["document_count"], 2)

    def test_invalid_jsonl_record_is_retained_with_a_warning(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "records-with-error.jsonl"
            path.write_text(
                '{"name":"valid-jsonl-sentinel"}\nnot-jsonl-secret\n',
                encoding="utf-8",
            )

            documents, report = self._load(tmpdir, path.name)

        self.assertEqual(len(documents), 2)
        self.assertIn("not-jsonl-secret", "\n".join(document.text for document in documents))
        self.assertIn("Invalid JSONL record", report[0]["warning"])
        self.assertEqual(report[0]["status"], "loaded")

    def test_csv_preserves_headers_rows_and_raw_table(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "people.csv"
            raw = "name,department\nAda,Research\nGrace,Engineering\n"
            path.write_text(raw, encoding="utf-8")

            documents, report = self._load(tmpdir, path.name)

        self.assertEqual(len(documents), 1)
        text = documents[0].text
        self.assertIn("Columns: name, department", text)
        self.assertIn("name is Ada", text)
        self.assertIn("department is Engineering", text)
        self.assertIn("Original table:", text)
        self.assertIn("name,department", text)
        self.assertIn("Ada,Research", text)
        self.assertIn("Grace,Engineering", text)
        self.assertEqual(report[0]["extracted_characters"], len(text))

    def test_html_reads_whole_document_without_script_or_style_text(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "page.html"
            path.write_text(
                "<html><head><style>style-secret</style></head>"
                "<body><h1>Visible HTML sentinel</h1>"
                "<script>script-secret</script><p>Body sentinel</p></body></html>",
                encoding="utf-8",
            )

            documents, report = self._load(tmpdir, path.name)

        text = documents[0].text
        self.assertIn("Visible HTML sentinel", text)
        self.assertIn("Body sentinel", text)
        self.assertNotIn("script-secret", text)
        self.assertNotIn("style-secret", text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_rtf_extracts_visible_text(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "message.rtf"
            path.write_text(
                r"{\rtf1\ansi Visible RTF sentinel\par}",
                encoding="ascii",
            )

            documents, report = self._load(tmpdir, path.name)

        self.assertIn("Visible RTF sentinel", documents[0].text)
        self.assertNotIn("rtf1", documents[0].text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_xml_extracts_element_text_safely(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "data.xml"
            path.write_text(
                "<root><title>XML sentinel</title><value>42</value></root>",
                encoding="utf-8",
            )

            documents, report = self._load(tmpdir, path.name)

        self.assertIn("XML sentinel", documents[0].text)
        self.assertIn("42", documents[0].text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_eml_extracts_plain_text_and_ignores_html_script(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "message.eml"
            path.write_bytes(
                b"From: sender@example.test\r\n"
                b"To: receiver@example.test\r\n"
                b"Subject: EML sentinel\r\n"
                b"Content-Type: text/plain; charset=utf-8\r\n\r\n"
                b"Email body sentinel\r\n"
            )

            documents, report = self._load(tmpdir, path.name)

        text = documents[0].text
        self.assertIn("Email body sentinel", text)
        self.assertIn("EML sentinel", text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_mhtml_extracts_mime_text(self):
        boundary = "docmind-boundary"
        content = (
            f"From: sender@example.test\r\n"
            f"Subject: MHTML sentinel\r\n"
            f"Content-Type: multipart/alternative; boundary={boundary}\r\n\r\n"
            f"--{boundary}\r\n"
            f"Content-Type: text/plain; charset=utf-8\r\n\r\n"
            f"MHTML body sentinel\r\n"
            f"--{boundary}--\r\n"
        ).encode("utf-8")
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "saved.mhtml"
            path.write_bytes(content)

            documents, report = self._load(tmpdir, path.name)

        self.assertIn("MHTML body sentinel", documents[0].text)
        self.assertIn("MHTML sentinel", documents[0].text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_mbox_retains_each_message(self):
        content = (
            b"From sender@example.test Mon Jan  1 00:00:00 2024\n"
            b"From: sender@example.test\n"
            b"Subject: MBOX first sentinel\n"
            b"Content-Type: text/plain; charset=utf-8\n\n"
            b"First MBOX body sentinel\n\n"
            b"From other@example.test Tue Jan  2 00:00:00 2024\n"
            b"From: other@example.test\n"
            b"Subject: MBOX second sentinel\n"
            b"Content-Type: text/plain; charset=utf-8\n\n"
            b"Second MBOX body sentinel\n\n"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "mailbox.mbox"
            path.write_bytes(content)

            documents, report = self._load(tmpdir, path.name)

        self.assertEqual(len(documents), 2)
        combined = "\n".join(document.text for document in documents)
        self.assertIn("First MBOX body sentinel", combined)
        self.assertIn("Second MBOX body sentinel", combined)
        self.assertEqual(report[0]["status"], "loaded")
        self.assertEqual(report[0]["document_count"], 2)

    def test_odt_extracts_paragraph_text(self):
        from odf import text as odf_text
        from odf.opendocument import OpenDocumentText

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "document.odt"
            document = OpenDocumentText()
            document.text.addElement(odf_text.P(text="ODT sentinel paragraph"))
            document.save(str(path))

            documents, report = self._load(tmpdir, path.name)

        self.assertIn("ODT sentinel paragraph", documents[0].text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_xlsx_extracts_header_and_data_row(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sheet.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["label", "amount"])
            worksheet.append(["XLSX sentinel", 17])
            workbook.save(path)
            workbook.close()

            documents, report = self._load(tmpdir, path.name)

        text = documents[0].text
        self.assertIn("Columns: label, amount", text)
        self.assertIn("label is XLSX sentinel", text)
        self.assertIn("amount is 17", text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_pptx_extracts_slide_text(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "PPTX sentinel title"
            presentation.save(path)

            documents, report = self._load(tmpdir, path.name)

        self.assertIn("PPTX sentinel title", "\n".join(document.text for document in documents))
        self.assertEqual(report[0]["status"], "loaded")

    def test_ipynb_extracts_notebook_code(self):
        notebook = {
            "cells": [
                {
                    "id": "cell-1",
                    "cell_type": "code",
                    "execution_count": None,
                    "metadata": {},
                    "outputs": [],
                    "source": ["IPYNB sentinel = 1\n"],
                }
            ],
            "metadata": {},
            "nbformat": 4,
            "nbformat_minor": 5,
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "notebook.ipynb"
            path.write_text(json.dumps(notebook), encoding="utf-8")

            documents, report = self._load(tmpdir, path.name)

        self.assertIn("IPYNB sentinel", "\n".join(document.text for document in documents))
        self.assertEqual(report[0]["status"], "loaded")

    def test_parser_failure_in_mixed_batch_is_reported_per_file(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            text_path = Path(tmpdir) / "usable.txt"
            xml_path = Path(tmpdir) / "broken.xml"
            text_path.write_text("Usable parser-failure sentinel", encoding="utf-8")
            xml_path.write_text(
                "<root><unclosed>parser-failure-secret</unclosed>",
                encoding="utf-8",
            )

            documents, report = load_documents(
                str(tmpdir),
                input_files=[str(text_path), str(xml_path)],
                return_report=True,
            )

        self.assertEqual(len(documents), 1)
        self.assertIn("Usable parser-failure sentinel", documents[0].text)
        statuses = {entry["filename"]: entry for entry in report}
        self.assertEqual(statuses["usable.txt"]["status"], "loaded")
        self.assertEqual(statuses["broken.xml"]["status"], "skipped")
        self.assertNotIn("parser-failure-secret", json.dumps(report))

    def test_legacy_binary_formats_are_rejected_with_conversion_guidance(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "legacy.doc"
            ppt_path = Path(tmpdir) / "legacy.ppt"
            doc_path.write_bytes(b"\xd0\xcf\x11\xe0legacy-doc-secret")
            ppt_path.write_bytes(b"\xd0\xcf\x11\xe0legacy-ppt-secret")

            with self.assertRaisesRegex(ValueError, "Convert.*docx"):
                load_documents(str(tmpdir), input_files=[str(doc_path)])
            with self.assertRaisesRegex(ValueError, "Convert.*pptx"):
                load_documents(str(tmpdir), input_files=[str(ppt_path)])

    def test_report_is_stored_without_document_content(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "private.txt"
            path.write_text("private-report-secret", encoding="utf-8")
            state = {}
            with patch(
                "utils.llama_index.st",
                SimpleNamespace(session_state=state),
            ):
                documents, report = self._load(tmpdir, path.name)

        self.assertEqual(state["file_extraction_report"], report)
        self.assertEqual(state["extraction_report"], report)
        self.assertNotIn("private-report-secret", json.dumps(report))
        self.assertIn("private-report-secret", documents[0].text)

    def test_mixed_batch_reports_unsupported_file_without_dropping_supported_file(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            text_path = Path(tmpdir) / "usable.txt"
            legacy_path = Path(tmpdir) / "legacy.ppt"
            text_path.write_text("Usable text sentinel", encoding="utf-8")
            legacy_path.write_bytes(b"legacy-secret")

            documents, report = load_documents(
                str(tmpdir),
                input_files=[str(text_path), str(legacy_path)],
                return_report=True,
            )

        self.assertEqual(len(documents), 1)
        self.assertIn("Usable text sentinel", documents[0].text)
        statuses = {entry["filename"]: entry["status"] for entry in report}
        self.assertEqual(statuses["usable.txt"], "loaded")
        self.assertEqual(statuses["legacy.ppt"], "unsupported")
        self.assertNotIn("legacy-secret", json.dumps(report))


if __name__ == "__main__":
    unittest.main()
