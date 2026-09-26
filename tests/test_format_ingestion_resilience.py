import io
import json
import os
import stat
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document as WordDocument
from ebooklib import epub
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from utils.format_ingestion import (
    PARSER_CAPABILITIES,
    SUPPORTED_EXTENSIONS,
    ArchiveLimits,
    IngestionResourceError,
    UnsafeArchiveError,
    archive_limits,
    build_file_extractors,
    validate_zip_archive,
)
from utils.llama_index import load_documents
from utils.rag_pipeline import validate_ingested_documents


class FormatFixtureHelpers:
    @staticmethod
    def png_bytes():
        output = io.BytesIO()
        Image.new("RGB", (3, 3), (20, 100, 180)).save(output, format="PNG")
        return output.getvalue()

    @staticmethod
    def pdf_bytes(text):
        escaped = (
            text.replace("\\", "\\\\")
            .replace("(", "\\(")
            .replace(")", "\\)")
            .encode("ascii")
        )
        stream = b"BT /F1 18 Tf 72 700 Td (" + escaped + b") Tj ET\n"
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length "
            + str(len(stream)).encode("ascii")
            + b" >>\nstream\n"
            + stream
            + b"endstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = [0]
        for number, value in enumerate(objects, start=1):
            offsets.append(len(payload))
            payload.extend(f"{number} 0 obj\n".encode("ascii"))
            payload.extend(value)
            payload.extend(b"\nendobj\n")
        xref_offset = len(payload)
        payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
        payload.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        payload.extend(
            (
                f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
                f"startxref\n{xref_offset}\n%%EOF\n"
            ).encode("ascii")
        )
        return bytes(payload)

    @staticmethod
    def blank_pdf(path):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        with path.open("wb") as handle:
            writer.write(handle)

    @staticmethod
    def make_epub(path):
        book = epub.EpubBook()
        book.set_identifier("docmind-resilience")
        book.set_title("Resilience fixture")
        book.set_language("en")
        first = epub.EpubHtml(
            uid="chapter-first",
            file_name="first.xhtml",
            lang="en",
            title="First",
        )
        first.content = "<html><body><h1>EPUB first sentinel</h1></body></html>"
        second = epub.EpubHtml(
            uid="chapter-second",
            file_name="second.xhtml",
            lang="en",
            title="Second",
        )
        second.content = "<html><body><p>EPUB second sentinel</p></body></html>"
        book.add_item(first)
        book.add_item(second)
        book.spine = ["chapter-first", "chapter-second"]
        book.toc = (
            epub.Link("first.xhtml", "First", "chapter-first"),
            epub.Link("second.xhtml", "Second", "chapter-second"),
        )
        ncx = epub.EpubNcx()
        ncx.content = (
            b"<?xml version='1.0' encoding='utf-8'?>"
            b"<ncx xmlns='http://www.daisy.org/z3986/2005/ncx/' version='2005-1'>"
            b"<navMap><navPoint id='first'><navLabel><text>First</text></navLabel>"
            b"<content src='first.xhtml'/></navPoint>"
            b"<navPoint id='second'><navLabel><text>Second</text></navLabel>"
            b"<content src='second.xhtml'/></navPoint></navMap></ncx>"
        )
        book.add_item(ncx)
        epub.EpubWriter(str(path), book).write()

    @staticmethod
    def make_docx(path, text="DOCX paragraph sentinel"):
        document = WordDocument()
        document.add_paragraph(text)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "DOCX table left"
        table.cell(0, 1).text = "DOCX table right"
        document.save(path)

    @staticmethod
    def make_odt(path):
        from odf import text as odf_text
        from odf.opendocument import OpenDocumentText

        document = OpenDocumentText()
        document.text.addElement(odf_text.P(text="ODT generated sentinel"))
        document.save(str(path))

    @staticmethod
    def make_xlsx(path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["label", "amount"])
        sheet.append(["XLSX generated sentinel", 19])
        workbook.save(path)
        workbook.close()

    @staticmethod
    def make_pptx(path, image_only=False):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        if image_only:
            slide.shapes.add_picture(
                io.BytesIO(FormatFixtureHelpers.png_bytes()),
                Inches(1),
                Inches(1),
                width=Inches(1),
            )
        else:
            slide.shapes.title.text = "PPTX generated title"
            table_shape = slide.shapes.add_table(
                1,
                2,
                Inches(1),
                Inches(2),
                Inches(4),
                Inches(1),
            )
            table_shape.table.cell(0, 0).text = "PPTX table left"
            table_shape.table.cell(0, 1).text = "PPTX table right"
            slide.notes_slide.notes_text_frame.text = "PPTX generated notes"
        presentation.save(path)

    @staticmethod
    def write_zip(path, entries, compression=zipfile.ZIP_DEFLATED):
        with zipfile.ZipFile(path, "w", compression=compression) as archive:
            for name, value in entries:
                archive.writestr(name, value)


class RealGeneratedFixtureTests(unittest.TestCase):
    def _load(self, directory, path):
        return load_documents(
            str(directory),
            input_files=[str(path)],
            return_report=True,
            store_report=False,
        )

    def test_docx_extracts_paragraph_and_table_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.docx"
            FormatFixtureHelpers.make_docx(path)
            documents, report = self._load(directory, path)
        text = "\n".join(document.text for document in documents)
        self.assertIn("DOCX paragraph sentinel", text)
        self.assertIn("DOCX table left", text)
        self.assertIn("DOCX table right", text)
        self.assertEqual(report[0]["status"], "loaded")
        self.assertEqual(report[0]["document_count"], 1)

    def test_epub_follows_spine_body_order(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.epub"
            FormatFixtureHelpers.make_epub(path)
            documents, report = self._load(directory, path)
        text = "\n".join(document.text for document in documents)
        self.assertLess(
            text.index("EPUB first sentinel"), text.index("EPUB second sentinel")
        )
        self.assertEqual(report[0]["status"], "loaded")

    def test_xlsx_extracts_table_values(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.xlsx"
            FormatFixtureHelpers.make_xlsx(path)
            documents, report = self._load(directory, path)
        text = "\n".join(document.text for document in documents)
        self.assertIn("Columns: label, amount", text)
        self.assertIn("label is XLSX generated sentinel", text)
        self.assertIn("amount is 19", text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_pptx_extracts_slide_table_and_notes(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.pptx"
            FormatFixtureHelpers.make_pptx(path)
            documents, report = self._load(directory, path)
        text = "\n".join(document.text for document in documents)
        self.assertIn("PPTX generated title", text)
        self.assertIn("PPTX table left", text)
        self.assertIn("PPTX generated notes", text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_odt_extracts_generated_paragraph(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.odt"
            FormatFixtureHelpers.make_odt(path)
            documents, report = self._load(directory, path)
        self.assertIn("ODT generated sentinel", documents[0].text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_pdf_with_text_layer_is_indexed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.pdf"
            path.write_bytes(FormatFixtureHelpers.pdf_bytes("PDF text layer sentinel"))
            documents, report = self._load(directory, path)
        self.assertIn("PDF text layer sentinel", documents[0].text)
        self.assertEqual(report[0]["status"], "loaded")

    def test_reader_mapping_retains_all_zip_readers(self):
        readers = build_file_extractors()
        self.assertEqual(
            {
                type(readers[extension]).__name__
                for extension in (".docx", ".pptx", ".xlsx", ".odt", ".epub")
            },
            {"DocxReader", "PptxReader", "XLSXReader", "ODTReader", "EpubReader"},
        )


class MalformedAndEmptyBatchTests(unittest.TestCase):
    def test_malformed_and_empty_supported_families_do_not_drop_valid_file_or_leak(
        self,
    ):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "good.txt").write_text(
                "valid mixed-batch sentinel", encoding="utf-8"
            )
            fixtures = {
                "empty.txt": b"",
                "bad-encoding.txt": b"\xffencoding-secret",
                "bad.json": b'{"secret":',
                "bad.jsonl": b'{"ok":1}\njsonl-secret',
                "empty.csv": b"",
                "bad.csv": b'"unterminated-secret',
                "empty.tsv": b"",
                "bad.tsv": b'"unterminated-secret',
                "bad.xml": b"<root><secret>truncated",
                "bad.htm": b"<html><unclosed>htm-secret",
                "bad.html": b"<html><unclosed>html-secret",
                "bad.md": b"",
                "bad.markdown": b"\xffmarkdown-secret",
                "bad.ipynb": b"{not-notebook-secret",
                "bad.eml": b"Subject: broken-secret\r\nContent-Type: multipart/mixed; boundary=x\r\n\r\n--x\r\n",
                "bad.mhtml": b"Content-Type: multipart/related; boundary=x\r\n\r\n--x\r\n",
                "bad.mbox": b"",
                "bad.rtf": b"{\\rtf",
                "bad.odt": b"PK\x03\x04odt-secret",
                "bad.docx": b"PK\x03\x04docx-secret",
                "bad.pptx": b"PK\x03\x04pptx-secret",
                "bad.xlsx": b"PK\x03\x04xlsx-secret",
                "bad.epub": b"PK\x03\x04epub-secret",
                "bad.xls": b"\x00xls-secret",
                "bad.pdf": b"%PDF-1.7\npdf-secret",
                "bad.msg": b"\xd0\xcf\x11\xe0msg-secret",
                "legacy.doc": b"\xd0\xcf\x11\xe0legacy-doc-secret",
                "legacy.ppt": b"\xd0\xcf\x11\xe0legacy-ppt-secret",
            }
            for name, payload in fixtures.items():
                (root / name).write_bytes(payload)
            input_files = [str(root / name) for name in ["good.txt", *fixtures]]
            with patch("utils.llama_index.logs.log") as log:
                documents, report = load_documents(
                    str(root),
                    input_files=input_files,
                    return_report=True,
                    store_report=False,
                )
        self.assertTrue(
            any("valid mixed-batch sentinel" in document.text for document in documents)
        )
        self.assertEqual(
            {entry["filename"] for entry in report}, {"good.txt", *fixtures}
        )
        self.assertEqual(
            next(entry for entry in report if entry["filename"] == "good.txt")[
                "status"
            ],
            "loaded",
        )
        self.assertEqual(
            next(entry for entry in report if entry["filename"] == "legacy.doc")[
                "status"
            ],
            "unsupported",
        )
        self.assertEqual(
            next(entry for entry in report if entry["filename"] == "legacy.ppt")[
                "status"
            ],
            "unsupported",
        )
        serialized = json.dumps(report)
        for secret in ("secret", "truncated", "unterminated"):
            self.assertNotIn(secret, serialized)
        self.assertNotIn("secret", str(log.mock_calls))
        statuses = {entry["filename"]: entry["status"] for entry in report}
        self.assertIn("bad.json", statuses)
        self.assertIn("bad.docx", statuses)
        self.assertIn("bad.pdf", statuses)
        self.assertTrue(
            all(
                status in {"loaded", "skipped", "unsupported"}
                for status in statuses.values()
            )
        )

    def test_invalid_jsonl_is_truthfully_loaded_with_a_bounded_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "records.jsonl"
            path.write_text('{"ok":1}\nnot-json-secret\n', encoding="utf-8")
            documents, report = load_documents(
                directory,
                input_files=[str(path)],
                return_report=True,
                store_report=False,
            )
        self.assertEqual(len(documents), 2)
        self.assertEqual(report[0]["status"], "loaded")
        self.assertIn("Invalid JSONL record", report[0]["warning"])
        self.assertNotIn("not-json-secret", json.dumps(report))

    def test_empty_legacy_formats_do_not_claim_support(self):
        with tempfile.TemporaryDirectory() as directory:
            doc = Path(directory) / "empty.doc"
            ppt = Path(directory) / "empty.ppt"
            doc.write_bytes(b"")
            ppt.write_bytes(b"")
            with self.assertRaises(ValueError):
                load_documents(directory, input_files=[str(doc)], store_report=False)
            with self.assertRaises(ValueError):
                load_documents(directory, input_files=[str(ppt)], store_report=False)


class ArchiveHardeningTests(unittest.TestCase):
    def _limits(self, **overrides):
        values = {
            "max_members": 20,
            "max_total_uncompressed_bytes": 10_000,
            "max_entry_uncompressed_bytes": 5_000,
            "max_compression_ratio": 500,
            "max_compressed_bytes": 100_000,
        }
        values.update(overrides)
        return ArchiveLimits(**values)

    def test_archive_limits_have_environment_overrides(self):
        with patch.dict(
            os.environ,
            {
                "DOCMIND_ZIP_MAX_MEMBERS": "7",
                "DOCMIND_ZIP_MAX_TOTAL_BYTES": "7000",
                "DOCMIND_ZIP_MAX_ENTRY_BYTES": "600",
                "DOCMIND_ZIP_MAX_COMPRESSION_RATIO": "2.5",
            },
            clear=False,
        ):
            limits = archive_limits()
        self.assertEqual(limits.max_members, 7)
        self.assertEqual(limits.max_total_uncompressed_bytes, 7000)
        self.assertEqual(limits.max_entry_uncompressed_bytes, 600)
        self.assertEqual(limits.max_compression_ratio, 2.5)

    def test_many_members_and_total_and_entry_limits_reject(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "many.zip"
            FormatFixtureHelpers.write_zip(
                path, [(f"x{index}", b"x") for index in range(3)]
            )
            with self.assertRaises(IngestionResourceError):
                validate_zip_archive(path, self._limits(max_members=2))
            with self.assertRaises(IngestionResourceError):
                validate_zip_archive(path, self._limits(max_total_uncompressed_bytes=2))
            with self.assertRaises(IngestionResourceError):
                validate_zip_archive(path, self._limits(max_entry_uncompressed_bytes=0))

    def test_high_ratio_and_zero_compression_ratio_reject(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ratio.zip"
            FormatFixtureHelpers.write_zip(path, [("large.txt", b"A" * 20_000)])
            with self.assertRaises(IngestionResourceError):
                validate_zip_archive(path, self._limits(max_compression_ratio=2))

    def test_traversal_absolute_symlink_special_encrypted_and_truncated_reject(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traversal = root / "traversal.zip"
            FormatFixtureHelpers.write_zip(traversal, [("../outside.txt", b"x")])
            with self.assertRaises(UnsafeArchiveError):
                validate_zip_archive(traversal, self._limits())
            absolute = root / "absolute.zip"
            FormatFixtureHelpers.write_zip(absolute, [("/outside.txt", b"x")])
            with self.assertRaises(UnsafeArchiveError):
                validate_zip_archive(absolute, self._limits())
            symlink = root / "symlink.zip"
            with zipfile.ZipFile(symlink, "w") as archive:
                info = zipfile.ZipInfo("link")
                info.create_system = 3
                info.external_attr = (stat.S_IFLNK | 0o777) << 16
                archive.writestr(info, b"target")
            with self.assertRaises(UnsafeArchiveError):
                validate_zip_archive(symlink, self._limits())
            special = root / "special.zip"
            with zipfile.ZipFile(special, "w") as archive:
                info = zipfile.ZipInfo("pipe")
                info.create_system = 3
                info.external_attr = (stat.S_IFIFO | 0o600) << 16
                archive.writestr(info, b"x")
            with self.assertRaises(UnsafeArchiveError):
                validate_zip_archive(special, self._limits())
            encrypted = root / "encrypted.zip"
            FormatFixtureHelpers.write_zip(encrypted, [("secret.txt", b"x")])
            encrypted_bytes = bytearray(encrypted.read_bytes())
            for signature, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
                start = 0
                while True:
                    position = encrypted_bytes.find(signature, start)
                    if position < 0:
                        break
                    flag_offset = position + offset
                    flags = int.from_bytes(
                        encrypted_bytes[flag_offset : flag_offset + 2], "little"
                    )
                    encrypted_bytes[flag_offset : flag_offset + 2] = (
                        flags | 1
                    ).to_bytes(2, "little")
                    start = position + 4
            encrypted.write_bytes(encrypted_bytes)
            with self.assertRaises(UnsafeArchiveError):
                validate_zip_archive(encrypted, self._limits())
            truncated = root / "truncated.zip"
            truncated.write_bytes(b"PK\x03\x04truncated")
            with self.assertRaises(UnsafeArchiveError):
                validate_zip_archive(truncated, self._limits())

    def test_unsafe_archive_is_rejected_before_docx_parser_is_called(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            unsafe = root / "unsafe.docx"
            FormatFixtureHelpers.write_zip(unsafe, [("../outside.txt", b"x")])
            good = root / "good.txt"
            good.write_text("survives archive rejection", encoding="utf-8")
            with patch(
                "utils.format_ingestion.DocxReader.load_data",
                side_effect=AssertionError("third-party parser was called"),
            ) as parser:
                documents, report = load_documents(
                    directory,
                    input_files=[str(unsafe), str(good)],
                    return_report=True,
                    store_report=False,
                )
        parser.assert_not_called()
        self.assertIn("survives archive rejection", documents[0].text)
        entry = next(item for item in report if item["filename"] == "unsafe.docx")
        self.assertEqual(entry["status"], "skipped")
        self.assertEqual(entry["error_category"], "unsafe_archive")


class TextContainerLimitTests(unittest.TestCase):
    def _mixed_load(self, directory, path):
        root = Path(directory)
        good = root / "limit-good.txt"
        good.write_text("limit valid sentinel", encoding="utf-8")
        return load_documents(
            directory,
            input_files=[str(good), str(path)],
            return_report=True,
            store_report=False,
        )

    def test_jsonl_record_count_and_deep_json_are_rejected_before_batch_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            jsonl = root / "large.jsonl"
            jsonl.write_text(
                "".join(f'{{"record":{index}}}\n' for index in range(4)),
                encoding="utf-8",
            )
            with patch.dict(os.environ, {"DOCMIND_MAX_RECORDS_PER_FILE": "3"}):
                documents, report = self._mixed_load(directory, jsonl)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            entry = next(item for item in report if item["filename"] == "large.jsonl")
            self.assertEqual(entry["status"], "skipped")
            self.assertIn("record count", entry["error"])
            deep = root / "deep.json"
            deep.write_text("[" * 6 + "0" + "]" * 6, encoding="utf-8")
            with patch.dict(os.environ, {"DOCMIND_MAX_JSON_DEPTH": "3"}):
                documents, report = self._mixed_load(directory, deep)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            entry = next(item for item in report if item["filename"] == "deep.json")
            self.assertEqual(entry["status"], "skipped")
            self.assertIn("nesting", entry["error"])
            array = root / "large.json"
            array.write_text(
                json.dumps([{"record": index} for index in range(4)]), encoding="utf-8"
            )
            with patch.dict(os.environ, {"DOCMIND_MAX_RECORDS_PER_FILE": "3"}):
                documents, report = self._mixed_load(directory, array)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            entry = next(item for item in report if item["filename"] == "large.json")
            self.assertEqual(entry["status"], "skipped")
            self.assertIn("record count", entry["error"])

    def test_large_mbox_and_mhtml_message_or_mime_counts_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            messages = []
            for index in range(4):
                messages.extend(
                    [
                        f"From sender{index}@example.test Mon Jan  1 00:00:00 2024\n".encode(),
                        b"From: sender@example.test\n",
                        b"Subject: message sentinel\n",
                        b"Content-Type: text/plain; charset=utf-8\n\n",
                        f"body {index}\n\n".encode(),
                    ]
                )
            mbox = root / "large.mbox"
            mbox.write_bytes(b"".join(messages))
            with patch.dict(os.environ, {"DOCMIND_MAX_INGESTED_DOCUMENTS": "3"}):
                documents, report = self._mixed_load(directory, mbox)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            entry = next(item for item in report if item["filename"] == "large.mbox")
            self.assertEqual(entry["status"], "skipped")
            self.assertIn("message count", entry["error"])
            mhtml = root / "large.mhtml"
            boundary = "limit-boundary"
            parts = []
            for index in range(4):
                parts.extend(
                    [
                        f"--{boundary}\r\n".encode(),
                        b"Content-Type: text/plain; charset=utf-8\r\n\r\n",
                        f"part {index}\r\n".encode(),
                    ]
                )
            mhtml.write_bytes(
                b"Content-Type: multipart/mixed; boundary="
                + boundary.encode()
                + b"\r\n\r\n"
                + b"".join(parts)
                + f"--{boundary}--\r\n".encode()
            )
            with patch.dict(os.environ, {"DOCMIND_MAX_MIME_PARTS": "2"}):
                documents, report = self._mixed_load(directory, mhtml)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            entry = next(item for item in report if item["filename"] == "large.mhtml")
            self.assertEqual(entry["status"], "skipped")
            self.assertIn("MIME part count", entry["error"])
            eml = root / "large.eml"
            eml.write_bytes(mhtml.read_bytes())
            with patch.dict(os.environ, {"DOCMIND_MAX_MIME_PARTS": "2"}):
                documents, report = self._mixed_load(directory, eml)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            entry = next(item for item in report if item["filename"] == "large.eml")
            self.assertEqual(entry["status"], "skipped")
            self.assertIn("MIME part count", entry["error"])

    def test_extracted_character_limit_applies_to_xml_and_html(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            xml = root / "huge.xml"
            xml.write_text("<root>" + "x" * 200 + "</root>", encoding="utf-8")
            with patch.dict(os.environ, {"DOCMIND_MAX_EXTRACTED_CHARS_PER_FILE": "20"}):
                documents, report = self._mixed_load(directory, xml)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            self.assertIn(
                "character limit",
                next(item for item in report if item["filename"] == "huge.xml")[
                    "error"
                ],
            )
            html = root / "huge.html"
            html.write_text("<p>" + "y" * 200 + "</p>", encoding="utf-8")
            with patch.dict(os.environ, {"DOCMIND_MAX_EXTRACTED_CHARS_PER_FILE": "20"}):
                documents, report = self._mixed_load(directory, html)
            self.assertTrue(
                any("limit valid sentinel" in document.text for document in documents)
            )
            self.assertIn(
                "character limit",
                next(item for item in report if item["filename"] == "huge.html")[
                    "error"
                ],
            )

    def test_document_and_total_text_boundaries_are_inclusive(self):
        from llama_index.core.schema import Document

        exact_documents = [Document(text="x") for _ in range(300)]
        self.assertIsNone(validate_ingested_documents(exact_documents))
        with self.assertRaises(ValueError):
            validate_ingested_documents([*exact_documents, Document(text="x")])
        exact_text = Document(text="x" * (4 * 1024 * 1024))
        self.assertIsNone(validate_ingested_documents([exact_text]))
        with self.assertRaises(ValueError):
            validate_ingested_documents([Document(text="x" * (4 * 1024 * 1024 + 1))])


class OcrAdjacentTests(unittest.TestCase):
    def _load_with_valid_text(self, directory, image_path):
        valid = Path(directory) / "valid.txt"
        valid.write_text("valid OCR-batch sentinel", encoding="utf-8")
        return load_documents(
            directory,
            input_files=[str(valid), str(image_path)],
            return_report=True,
            store_report=False,
        )

    def test_blank_pdf_is_skipped_with_explicit_ocr_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "blank.pdf"
            FormatFixtureHelpers.blank_pdf(path)
            documents, report = self._load_with_valid_text(directory, path)
        self.assertTrue(
            any("valid OCR-batch sentinel" in document.text for document in documents)
        )
        self.assertFalse(
            any("blank" in document.text.lower() for document in documents)
        )
        entry = next(item for item in report if item["filename"] == "blank.pdf")
        self.assertEqual(entry["status"], "skipped")
        self.assertRegex(entry["warning"], r"(?i)ocr|conversion")

    def test_image_only_pdf_is_skipped_with_explicit_ocr_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "image-only.pdf"
            Image.new("RGB", (12, 12), (200, 30, 30)).save(path, format="PDF")
            documents, report = self._load_with_valid_text(directory, path)
        self.assertTrue(
            any("valid OCR-batch sentinel" in document.text for document in documents)
        )
        entry = next(item for item in report if item["filename"] == "image-only.pdf")
        self.assertEqual(entry["status"], "skipped")
        self.assertRegex(entry["warning"], r"(?i)ocr|conversion")

    def test_image_only_docx_and_pptx_are_not_falsely_indexed(self):
        with tempfile.TemporaryDirectory() as directory:
            docx = Path(directory) / "image-only.docx"
            document = WordDocument()
            document.add_picture(io.BytesIO(FormatFixtureHelpers.png_bytes()))
            document.save(docx)
            documents, report = self._load_with_valid_text(directory, docx)
            self.assertTrue(
                any(
                    "valid OCR-batch sentinel" in document.text
                    for document in documents
                )
            )
            entry = next(
                item for item in report if item["filename"] == "image-only.docx"
            )
            self.assertEqual(entry["status"], "skipped")
            self.assertRegex(entry["warning"], r"(?i)ocr|conversion")
            pptx = Path(directory) / "image-only.pptx"
            FormatFixtureHelpers.make_pptx(pptx, image_only=True)
            documents, report = self._load_with_valid_text(directory, pptx)
            self.assertTrue(
                any(
                    "valid OCR-batch sentinel" in document.text
                    for document in documents
                )
            )
            entry = next(
                item for item in report if item["filename"] == "image-only.pptx"
            )
            self.assertEqual(entry["status"], "skipped")
            self.assertRegex(entry["warning"], r"(?i)ocr|conversion")
            self.assertFalse(
                any("image-only" in document.text for document in documents)
            )

    def test_capability_matrix_keeps_legacy_and_ocr_claims_truthful(self):
        self.assertEqual(
            PARSER_CAPABILITIES["doc"].category, "unsupported_without_converter"
        )
        self.assertEqual(
            PARSER_CAPABILITIES["ppt"].category, "unsupported_without_converter"
        )
        self.assertNotIn("OCR", PARSER_CAPABILITIES["doc"].note)
        self.assertNotIn("OCR", PARSER_CAPABILITIES["ppt"].note)
        self.assertIn("external OCR engine", PARSER_CAPABILITIES["pdf"].note)
        self.assertIn("external OCR", PARSER_CAPABILITIES["pptx"].note)
        self.assertNotIn("OCR", PARSER_CAPABILITIES["pdf"].dependency)
        self.assertEqual(len(SUPPORTED_EXTENSIONS), 25)

    def test_reports_retain_source_and_generation_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "metadata.txt"
            path.write_text("metadata sentinel", encoding="utf-8")
            state = {}
            with patch(
                "utils.llama_index.st",
                SimpleNamespace(session_state=state),
            ):
                _, report = load_documents(
                    directory,
                    input_files=[str(path)],
                    return_report=True,
                    source_id="source-123",
                    index_generation=7,
                )
        self.assertEqual(report[0]["source_id"], "source-123")
        self.assertEqual(report[0]["index_generation"], 7)
        self.assertEqual(state["file_extraction_report"][0]["source_id"], "source-123")
        self.assertEqual(state["extraction_report"][0]["index_generation"], 7)


if __name__ == "__main__":

    unittest.main()
