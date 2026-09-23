import io
from datetime import datetime

import docx
import streamlit as st

import utils.ollama as ollama
import utils.r2r as r2r
from components.page_state import default_chat_model
from utils.browser_settings import ensure_ollama_endpoint
from utils.endpoint_policy import normalize_provider_endpoint
from utils.provider_config import (
    OLLAMA,
    clear_model_catalog,
    get_chat_profile,
    get_embedding_profile,
    initialize_provider_state,
    normalize_provider_kind,
    profile_keys,
    provider_label,
)
from utils.source_state import normalize_chunk_settings


def _check_r2r_connection():
    """Probe the configured R2R server and record the result for the UI."""
    st.session_state["r2r_connection_ok"] = r2r.get_client().health()


def _clear_r2r_connection_check():
    st.session_state["r2r_connection_ok"] = None


def _style_to_prompt(style: str) -> str:
    """Return the system prompt for a given answer style preset."""
    base = (
        "You are DocMind AI, a helpful and accurate virtual assistant. "
        "When document context is provided, answer strictly from that context "
        "and do not invent information. Otherwise answer from your general "
        "knowledge. Be factual and conversational."
    )
    style_instructions = {
        "Concise": " Keep answers as short as possible — one or two sentences.",
        "Balanced (default)": " Keep answers concise but complete.",
        "Detailed": " Provide thorough explanations with context and examples.",
        "Bulleted": " Structure answers as bullet points for readability.",
        "Technical": " Use precise terminology; assume a technical audience.",
        "Simple / ELI5": " Explain simply, like you're talking to a 12-year-old. Avoid jargon.",
    }
    return base + style_instructions.get(style, style_instructions["Balanced (default)"])


BACKEND_PRESETS = {
    "Ollama": {
        "base_url": "http://localhost:11434",
        "chat_model": "",
        "embedding_model": "",
    },
    "OpenAI": {
        "base_url": "https://api.openai.com/v1",
        "chat_model": "gpt-4o-mini",
        "embedding_model": "text-embedding-3-small",
    },
    "LM Studio (Local AI)": {
        "base_url": "http://localhost:1234/v1",
        "chat_model": "local-model",
        "embedding_model": "",
    },
    "TabbyAPI": {
        "base_url": "http://localhost:5000/v1",
        "chat_model": "local-model",
        "embedding_model": "",
    },
    "OpenAI-compatible": {
        "base_url": "http://localhost:1234/v1",
        "chat_model": "local-model",
        "embedding_model": "",
    },
}
EMBEDDING_BACKENDS = list(BACKEND_PRESETS)


def _chat_profile(backend=None):
    return get_chat_profile(
        st.session_state,
        backend or st.session_state.get("llm_backend", "Ollama"),
    )


def _clear_profile_models(profile):
    keys = profile.get("keys") or {}
    models_key = keys.get("models")
    if models_key:
        st.session_state[models_key] = []
    clear_model_catalog(st.session_state, profile)


def _apply_backend_preset():
    preset = st.session_state.get("llm_backend", "Ollama")
    initialize_provider_state(st.session_state)
    details = BACKEND_PRESETS.get(preset, {})
    if preset == "Ollama":
        if not st.session_state.get("ollama_endpoint"):
            st.session_state["ollama_endpoint"] = details["base_url"]
        return
    profile = _chat_profile(preset)
    keys = profile["keys"]
    if not st.session_state.get(keys["base_url"]):
        st.session_state[keys["base_url"]] = details["base_url"]
    if not st.session_state.get(keys["model"]):
        st.session_state[keys["model"]] = details["chat_model"]
    st.session_state["active_chat_provider_kind"] = profile["provider_kind"]
    _clear_profile_models(profile)


def _on_chat_provider_change():
    _apply_backend_preset()


def _on_chat_url_change():
    backend = st.session_state.get("llm_backend", "Ollama")
    keys = profile_keys(backend)
    try:
        endpoint = normalize_provider_endpoint(
            st.session_state.get(keys["base_url"]),
            label=f"{provider_label(backend)} endpoint",
        )
    except ValueError as err:
        st.session_state["provider_endpoint_error"] = str(err)
        return
    st.session_state["provider_endpoint_error"] = None
    binding_key = f"{keys['api_key']}_endpoint"
    if st.session_state.get(keys["api_key"]) and st.session_state.get(binding_key) != endpoint:
        st.session_state[keys["api_key"]] = ""
    st.session_state[keys["base_url"]] = endpoint
    st.session_state[binding_key] = endpoint
    _clear_profile_models({"keys": keys})


def _on_chat_key_change():
    backend = st.session_state.get("llm_backend", "Ollama")
    keys = profile_keys(backend)
    st.session_state[f"{keys['api_key']}_endpoint"] = st.session_state.get(
        keys["base_url"]
    )
    _clear_profile_models({"keys": keys})


def _fetch_openai_models():
    profile = _chat_profile()
    try:
        catalog = ollama.get_openai_model_catalog(
            profile["base_url"],
            profile["api_key"],
            provider_kind=profile["provider_kind"],
        )
    except ValueError as err:
        st.session_state["provider_model_error"] = str(err)
        _clear_profile_models(profile)
        return
    keys = profile["keys"]
    st.session_state[keys["models"]] = catalog["chat"]
    st.session_state[f"{keys['models']}_endpoint"] = catalog["endpoint"]
    st.session_state[f"{keys['models']}_key_fingerprint"] = catalog["key_fingerprint"]
    st.session_state[f"{keys['models']}_provider"] = profile["provider_kind"]
    st.session_state["openai_models"] = catalog["all"]
    st.session_state["openai_models_endpoint"] = catalog["endpoint"]
    st.session_state["openai_models_key_fingerprint"] = catalog["key_fingerprint"]
    st.session_state["openai_models_provider"] = profile["provider_kind"]
    st.session_state["provider_model_error"] = None
    if profile["model"] not in catalog["chat"]:
        st.session_state[keys["model"]] = catalog["chat"][0] if catalog["chat"] else profile["model"]


def _refresh_models():
    try:
        ensure_ollama_endpoint(st.session_state)
    except ValueError as err:
        st.session_state["provider_endpoint_error"] = str(err)
        return
    initialize_provider_state(st.session_state)
    ollama.get_models()
    ollama.get_embedding_models()
    if st.session_state.get("selected_model") not in st.session_state["ollama_models"]:
        st.session_state["selected_model"] = default_chat_model(
            st.session_state["ollama_models"]
        )
    st.session_state["ollama_models_endpoint"] = st.session_state["ollama_endpoint"]
    st.session_state["ollama_embedding_models_endpoint"] = st.session_state["ollama_endpoint"]


def _on_embedding_backend_change():
    backend = st.session_state.get("embedding_backend", "Ollama")
    kind = normalize_provider_kind(backend)
    if kind == OLLAMA:
        default_url = BACKEND_PRESETS["Ollama"]["base_url"]
    else:
        profile = get_chat_profile(st.session_state, backend)
        default_url = profile["base_url"]
    st.session_state["embedding_base_url"] = default_url
    st.session_state["embedding_model"] = ""
    st.session_state["embedding_api_key"] = ""
    st.session_state["embedding_models"] = []
    st.session_state["embedding_models_endpoint"] = None
    st.session_state["embedding_models_key_fingerprint"] = None
    st.session_state["embedding_models_provider"] = None
    st.session_state["embedding_api_key_endpoint"] = None
    st.session_state["embedding_api_key_provider"] = None


def _on_embedding_url_change():
    try:
        endpoint = normalize_provider_endpoint(
            st.session_state.get("embedding_base_url"),
            label="Embedding endpoint",
        )
    except ValueError as err:
        st.session_state["embedding_endpoint_error"] = str(err)
        return
    st.session_state["embedding_endpoint_error"] = None
    binding = st.session_state.get("embedding_api_key_endpoint")
    if st.session_state.get("embedding_api_key") and binding != endpoint:
        st.session_state["embedding_api_key"] = ""
        st.session_state["embedding_api_key_provider"] = None
    st.session_state["embedding_base_url"] = endpoint
    st.session_state["embedding_api_key_endpoint"] = endpoint
    st.session_state["embedding_models"] = []
    st.session_state["embedding_models_endpoint"] = None
    st.session_state["embedding_models_key_fingerprint"] = None
    st.session_state["embedding_models_provider"] = None


def _on_embedding_key_change():
    st.session_state["embedding_api_key_endpoint"] = st.session_state.get(
        "embedding_base_url"
    )
    st.session_state["embedding_api_key_provider"] = normalize_provider_kind(
        st.session_state.get("embedding_backend", "Ollama")
    )
    st.session_state["embedding_models"] = []
    st.session_state["embedding_models_endpoint"] = None
    st.session_state["embedding_models_key_fingerprint"] = None
    st.session_state["embedding_models_provider"] = None


def _fetch_embedding_models():
    profile = get_embedding_profile(st.session_state)
    if profile["provider_kind"] == OLLAMA:
        st.session_state["embedding_base_url"] = profile["base_url"]
        ollama.get_embedding_models(profile["base_url"])
        st.session_state["embedding_models"] = list(
            st.session_state.get("ollama_embedding_models", [])
        )
        st.session_state["ollama_embedding_models_endpoint"] = profile["base_url"]
        return
    try:
        catalog = ollama.get_openai_model_catalog(
            profile["base_url"],
            profile["api_key"],
            provider_kind=profile["provider_kind"],
        )
    except ValueError as err:
        st.session_state["embedding_model_error"] = str(err)
        st.session_state["embedding_models"] = []
        return
    st.session_state["embedding_models"] = catalog["embedding"]
    st.session_state["embedding_model_error"] = None
    st.session_state["embedding_models_endpoint"] = catalog["endpoint"]
    st.session_state["embedding_models_key_fingerprint"] = catalog["key_fingerprint"]
    st.session_state["embedding_models_provider"] = profile["provider_kind"]


def _refresh_embedding_models():
    try:
        ensure_ollama_endpoint(st.session_state)
    except ValueError as err:
        st.session_state["provider_endpoint_error"] = str(err)
        return
    initialize_provider_state(st.session_state)
    _fetch_embedding_models()
    embedding = get_embedding_profile(st.session_state)
    if embedding["provider_kind"] == OLLAMA:
        st.session_state["ollama_embedding_models_endpoint"] = embedding["base_url"]


def _chat_history_signature(messages):
    """Return a hashable snapshot of the chat history for cache keying."""
    return tuple((m.get("role"), m.get("content")) for m in messages)


@st.cache_data(show_spinner=False)
def chat_history_docx(signature):
    """Serialize chat history to a Word document, in memory."""
    document = docx.Document()
    document.add_heading("DocMind AI — Chat History", level=0)
    document.add_paragraph(f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    for role, content in signature:
        paragraph = document.add_paragraph()
        label = "User" if role == "user" else "Assistant"
        paragraph.add_run(f"[{label}] ").bold = True
        paragraph.add_run(str(content))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def settings():
    try:
        normalize_chunk_settings(st.session_state)
    except ValueError:
        normalize_chunk_settings(st.session_state, strict=False)
    st.header("Settings")
    st.caption("Pick your models and answer style — everything else is optional.")

    initialize_provider_state(st.session_state)
    st.subheader("Chat")
    st.caption("The model that writes your answers.")
    with st.container(border=True):
        backend = st.selectbox(
            "Provider",
            options=list(BACKEND_PRESETS.keys()),
            key="llm_backend",
            on_change=_on_chat_provider_change,
            help="Ollama runs offline. OpenAI-compatible providers use their own endpoint and credentials.",
        )
        if backend == "Ollama":
            ollama_models = st.session_state.get("ollama_models", [])
            st.selectbox(
                "Chat Model",
                ollama_models,
                key="selected_model",
                disabled=len(ollama_models) == 0,
                placeholder="Select Chat Model" if ollama_models else "No Models Available",
            )
            st.button("Refresh Models", key="refresh_chat_models", on_click=_refresh_models)
            if not ollama_models:
                st.info("No models found. In a terminal run: `ollama pull qwen2.5:0.5b`")
            with st.expander("Connection", expanded=False):
                st.text_input(
                    "Ollama Endpoint",
                    key="ollama_endpoint",
                    placeholder="http://localhost:11434",
                    on_change=_refresh_models,
                    help="Use a loopback HTTP URL or an HTTPS URL.",
                )
        else:
            profile = _chat_profile(backend)
            keys = profile["keys"]
            st.text_input(
                "Server URL",
                key=keys["base_url"],
                placeholder=profile["base_url"],
                on_change=_on_chat_url_change,
                help="Use a loopback HTTP URL or an HTTPS URL; credentials and query strings are rejected.",
            )
            st.text_input(
                "API Key",
                key=keys["api_key"],
                type="password",
                on_change=_on_chat_key_change,
                help="Required for official OpenAI; optional for local servers.",
            )
            chat_models = list(st.session_state.get(keys["models"], []) or [])
            if chat_models:
                if profile["model"] not in chat_models:
                    chat_models.insert(0, profile["model"])
                st.selectbox(
                    "Chat Model",
                    chat_models,
                    key=keys["model"],
                )
            else:
                st.text_input(
                    "Chat Model",
                    key=keys["model"],
                    placeholder="local-model or a server model ID",
                )
            st.button(
                "Fetch Models from Server",
                key="fetch_openai_models",
                on_click=_fetch_openai_models,
            )
            if st.session_state.get("provider_model_error"):
                st.error(st.session_state["provider_model_error"])
            found = st.session_state.get("openai_models") or []
            if found:
                st.caption(
                    "Found: "
                    + ", ".join(found[:10])
                    + ("…" if len(found) > 10 else "")
                )
        if st.session_state.get("provider_endpoint_error"):
            st.error(st.session_state["provider_endpoint_error"])

    st.subheader("Document Search")
    st.caption("Embedding chat and document models can use independent provider profiles.")
    with st.container(border=True):
        embedding_backend = st.session_state.get("embedding_backend", "Ollama")
        if embedding_backend not in EMBEDDING_BACKENDS:
            embedding_backend = provider_label(embedding_backend)
            st.session_state["embedding_backend"] = embedding_backend
        embedding_backend = st.selectbox(
            "Embedding Provider",
            EMBEDDING_BACKENDS,
            key="embedding_backend",
            on_change=_on_embedding_backend_change,
        )
        if embedding_backend == "Ollama":
            st.text_input(
                "Embedding Server URL",
                key="embedding_base_url",
                placeholder="http://localhost:11434",
                on_change=_on_embedding_url_change,
            )
            embedding_models = list(st.session_state.get("embedding_models", []) or [])
            if not embedding_models:
                embedding_models = list(
                    st.session_state.get("ollama_embedding_models", []) or []
                )
            if embedding_models:
                current = st.session_state.get("embedding_model")
                if current not in embedding_models:
                    embedding_models.insert(0, current)
                st.selectbox(
                    "Embedding Model",
                    embedding_models,
                    key="embedding_model",
                )
            else:
                st.text_input(
                    "Embedding Model",
                    key="embedding_model",
                    placeholder="nomic-embed-text:latest",
                )
            st.button(
                "Fetch Embedding Models",
                key="fetch_embedding_models",
                on_click=_fetch_embedding_models,
            )
        else:
            st.text_input(
                "Embedding Server URL",
                key="embedding_base_url",
                placeholder="http://localhost:1234/v1",
                on_change=_on_embedding_url_change,
            )
            st.text_input(
                "Embedding API Key",
                key="embedding_api_key",
                type="password",
                on_change=_on_embedding_key_change,
                help="Leave empty for local servers; official OpenAI requires an explicit key.",
            )
            embedding_models = list(st.session_state.get("embedding_models", []) or [])
            if embedding_models:
                current = st.session_state.get("embedding_model")
                if current and current not in embedding_models:
                    embedding_models.insert(0, current)
                st.selectbox(
                    "Embedding Model",
                    embedding_models,
                    key="embedding_model",
                )
            else:
                st.text_input(
                    "Embedding Model",
                    key="embedding_model",
                    placeholder="text-embedding-3-small or a server model ID",
                )
            st.button(
                "Fetch Embedding Models",
                key="fetch_embedding_models",
                on_click=_fetch_embedding_models,
            )
        if st.session_state.get("embedding_endpoint_error"):
            st.error(st.session_state["embedding_endpoint_error"])
        if st.session_state.get("embedding_model_error"):
            st.error(st.session_state["embedding_model_error"])

    # ── 3. Answer Style ────────────────────────────────────────────
    st.subheader("Answer Style")
    st.caption("How the assistant should sound.")
    with st.container(border=True):
        style = st.radio(
            "Style",
            options=[
                "Concise",
                "Balanced (default)",
                "Detailed",
                "Bulleted",
                "Technical",
                "Simple / ELI5",
            ],
            index=1,
            horizontal=True,
            key="answer_style",
        )
        st.session_state["system_prompt"] = _style_to_prompt(style)
        # Keep chatbox quick selector in sync
        st.session_state["quick_answer_style"] = style
        with st.expander("Preview prompt", expanded=False):
            st.caption(_style_to_prompt(style))

    # ── 4. Everyday toggles ────────────────────────────────────────
    st.subheader("Preferences")
    with st.container(border=True):
        st.toggle(
            "Eco Mode — save power & stay cool",
            key="eco_mode",
            help="Fewer chunks, shorter answers, smaller batches. Turn on if the laptop gets hot or slow.",
        )
        if st.session_state.get("eco_mode"):
            st.caption("On: smaller batches · shorter answers · fewer chunks.")
        else:
            st.caption("Off: best quality. Turn on when the machine runs hot.")
        st.divider()
        st.toggle("Show advanced controls", key="advanced")
        st.caption("Reveals fine-tuning sliders and server options below.")

    # ── 5. Advanced (only when enabled) ─────────────────────────────
    if st.session_state.get("advanced"):
        with st.container(border=True):
            st.markdown("**Fine-tuning**")
            st.select_slider(
                "Sources per answer",
                options=[1, 2, 3, 4, 5, 6, 7, 8, 9, 10],
                value=st.session_state["top_k"],
                key="top_k",
                help="How many document chunks to include when answering.",
            )
            st.slider(
                "Candidate depth",
                min_value=1,
                max_value=50,
                value=int(st.session_state.get("candidate_depth", 10)),
                key="candidate_depth",
                help="Bounded depth for each independent retrieval pool.",
            )
            st.slider(
                "Vector candidates",
                min_value=1,
                max_value=50,
                value=int(st.session_state.get("vector_candidate_depth", 10)),
                key="vector_candidate_depth",
            )
            st.slider(
                "BM25 candidates",
                min_value=1,
                max_value=50,
                value=int(st.session_state.get("bm25_candidate_depth", 10)),
                key="bm25_candidate_depth",
            )
            st.slider(
                "Relevance threshold",
                min_value=0.0,
                max_value=1.0,
                step=0.05,
                value=st.session_state.get("similarity_cutoff", 0.3),
                help="Higher = stricter, only very relevant chunks.",
                key="similarity_cutoff",
            )
            st.slider(
                "Creativity",
                min_value=0.0,
                max_value=1.5,
                step=0.05,
                value=float(st.session_state.get("temperature", 0.4)),
                help="Lower = focused & factual, higher = more creative.",
                key="temperature",
            )
            st.text_input(
                "Chunk Size",
                key="chunk_size",
                value=st.session_state["chunk_size"],
                help="Tokens per document piece (~4 chars each). Smaller = more precise, more work.",
            )
            chunk_overlap_pct = st.slider(
                "Chunk Overlap",
                min_value=0,
                max_value=50,
                step=1,
                value=st.session_state.get("chunk_overlap_pct", 12),
                help="Overlap between chunks as % of chunk size. Keeps context across boundaries.",
                key="chunk_overlap_pct",
            )
            try:
                normalized = normalize_chunk_settings(
                    st.session_state,
                    chunk_size=st.session_state.get("chunk_size"),
                    chunk_overlap_pct=chunk_overlap_pct,
                )
            except ValueError as err:
                st.error(str(err))
                normalized = normalize_chunk_settings(
                    st.session_state,
                    strict=False,
                )
            st.caption(f"→ {normalized['chunk_overlap']} tokens overlap")

    # ── 6. R2R (collapsed by default) ──────────────────────────────
    with st.expander("External RAG server (R2R) — optional", expanded=False):
        st.caption("Use an external R2R server instead of local indexing. Most users can ignore this.")
        st.caption(
            f"Wire contract: R2R {r2r.R2R_SERVER_CONTRACT} / API {r2r.R2R_API_CONTRACT}. "
            "Live compatibility is not claimed without an integration test."
        )
        st.toggle(
            "Enable R2R",
            key="r2r_enabled",
            help="When on, files are indexed on the R2R server.",
        )
        st.text_input(
            "R2R URL",
            key="r2r_base_url",
            placeholder=r2r.DEFAULT_R2R_BASE_URL,
            on_change=_clear_r2r_connection_check,
        )
        st.text_input(
            "API Key (optional)",
            key="r2r_api_key",
            type="password",
        )
        st.button(
            "Test Connection",
            key="check_r2r_connection",
            on_click=_check_r2r_connection,
        )
        if st.session_state.get("r2r_connection_ok") is True:
            st.success("R2R server reachable.")
        elif st.session_state.get("r2r_connection_ok") is False:
            st.error("Could not reach R2R. Check the URL or disable R2R.")
        if st.session_state.get("r2r_enabled") and not st.session_state.get("r2r_document_ids"):
            st.info("R2R is on — upload files in Data Sources → Local Files.")

    # ── 7. Export ──────────────────────────────────────────────────
    st.subheader("Export")
    with st.container(border=True):
        st.caption("Save your conversation as a Word file.")
        st.download_button(
            label="Download chat (.docx)",
            data=chat_history_docx(_chat_history_signature(st.session_state["messages"])),
            file_name=f"docmind-chat-{datetime.now().strftime('%Y-%m-%d-%H-%M-%S')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
